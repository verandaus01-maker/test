"""
Export System
Multi-format data export functionality
"""

from typing import List, Dict, Any, Optional
import csv
import json
import xml.etree.ElementTree as ET
from io import StringIO, BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import xlsxwriter
import logging

logger = logging.getLogger(__name__)


class BaseExporter:
    """Base class for all exporters"""

    def __init__(self):
        self.created_at = datetime.utcnow()

    def export(self, leads: List[Dict[str, Any]], fields: Optional[List[str]] = None) -> bytes:
        """
        Export leads to specific format

        Args:
            leads: List of lead dictionaries
            fields: Optional list of fields to include

        Returns:
            Exported data as bytes
        """
        raise NotImplementedError


class CSVExporter(BaseExporter):
    """Export leads to CSV format"""

    def export(self, leads: List[Dict[str, Any]], fields: Optional[List[str]] = None) -> bytes:
        """
        Export to CSV

        Args:
            leads: List of leads
            fields: Fields to include

        Returns:
            CSV data as bytes
        """
        if not leads:
            return b""

        # Determine fields
        if fields is None:
            fields = list(leads[0].keys())

        # Create CSV
        output = StringIO()
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction='ignore')

        # Write header
        writer.writeheader()

        # Write data
        for lead in leads:
            # Convert complex types to strings
            row = {}
            for field in fields:
                value = lead.get(field)
                if isinstance(value, (list, dict)):
                    row[field] = json.dumps(value)
                elif isinstance(value, datetime):
                    row[field] = value.isoformat()
                else:
                    row[field] = value
            writer.writerow(row)

        return output.getvalue().encode('utf-8')


class ExcelExporter(BaseExporter):
    """Export leads to Excel format (.xlsx)"""

    def export(self, leads: List[Dict[str, Any]], fields: Optional[List[str]] = None) -> bytes:
        """
        Export to Excel

        Args:
            leads: List of leads
            fields: Fields to include

        Returns:
            Excel file as bytes
        """
        if not leads:
            return b""

        # Determine fields
        if fields is None:
            fields = list(leads[0].keys())

        # Create workbook
        output = BytesIO()
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Leads"

        # Header styling
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # Write headers
        for col_idx, field in enumerate(fields, 1):
            cell = sheet.cell(row=1, column=col_idx, value=field.replace('_', ' ').title())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # Write data
        for row_idx, lead in enumerate(leads, 2):
            for col_idx, field in enumerate(fields, 1):
                value = lead.get(field)

                # Convert complex types
                if isinstance(value, (list, dict)):
                    value = json.dumps(value)
                elif isinstance(value, datetime):
                    value = value.isoformat()

                sheet.cell(row=row_idx, column=col_idx, value=value)

        # Auto-adjust column widths
        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            sheet.column_dimensions[column_letter].width = adjusted_width

        # Save to bytes
        workbook.save(output)
        return output.getvalue()


class JSONExporter(BaseExporter):
    """Export leads to JSON format"""

    def export(self, leads: List[Dict[str, Any]], fields: Optional[List[str]] = None) -> bytes:
        """
        Export to JSON

        Args:
            leads: List of leads
            fields: Fields to include

        Returns:
            JSON data as bytes
        """
        if not leads:
            return b"[]"

        # Filter fields if specified
        if fields:
            filtered_leads = []
            for lead in leads:
                filtered_lead = {field: lead.get(field) for field in fields if field in lead}
                filtered_leads.append(filtered_lead)
        else:
            filtered_leads = leads

        # Convert datetime objects
        def datetime_converter(obj):
            if isinstance(obj, datetime):
                return obj.isoformat()
            raise TypeError(f"Type {type(obj)} not serializable")

        # Export to JSON
        json_data = json.dumps(
            {
                'exported_at': datetime.utcnow().isoformat(),
                'total_leads': len(filtered_leads),
                'leads': filtered_leads
            },
            indent=2,
            default=datetime_converter
        )

        return json_data.encode('utf-8')


class XMLExporter(BaseExporter):
    """Export leads to XML format"""

    def export(self, leads: List[Dict[str, Any]], fields: Optional[List[str]] = None) -> bytes:
        """
        Export to XML

        Args:
            leads: List of leads
            fields: Fields to include

        Returns:
            XML data as bytes
        """
        if not leads:
            return b'<?xml version="1.0" encoding="UTF-8"?><leads></leads>'

        # Create root element
        root = ET.Element('leads')
        root.set('exported_at', datetime.utcnow().isoformat())
        root.set('total', str(len(leads)))

        # Add each lead
        for lead in leads:
            lead_element = ET.SubElement(root, 'lead')

            # Determine fields
            lead_fields = fields if fields else lead.keys()

            for field in lead_fields:
                value = lead.get(field)

                if value is not None:
                    field_element = ET.SubElement(lead_element, field)

                    # Convert value to string
                    if isinstance(value, (list, dict)):
                        field_element.text = json.dumps(value)
                    elif isinstance(value, datetime):
                        field_element.text = value.isoformat()
                    else:
                        field_element.text = str(value)

        # Convert to string
        xml_string = ET.tostring(root, encoding='utf-8', method='xml')
        return xml_string


class VCFExporter(BaseExporter):
    """Export leads to VCF (vCard) format for contact management"""

    def export(self, leads: List[Dict[str, Any]], fields: Optional[List[str]] = None) -> bytes:
        """
        Export to VCF format

        Args:
            leads: List of leads
            fields: Not used for VCF

        Returns:
            VCF data as bytes
        """
        vcf_lines = []

        for lead in leads:
            vcf_lines.append("BEGIN:VCARD")
            vcf_lines.append("VERSION:3.0")

            # Company name
            if lead.get('company_name'):
                vcf_lines.append(f"ORG:{lead['company_name']}")

            # Contact name
            if lead.get('contact_name'):
                vcf_lines.append(f"FN:{lead['contact_name']}")

            # Title
            if lead.get('contact_title'):
                vcf_lines.append(f"TITLE:{lead['contact_title']}")

            # Email
            if lead.get('email'):
                vcf_lines.append(f"EMAIL;TYPE=WORK:{lead['email']}")
            elif lead.get('contact_email'):
                vcf_lines.append(f"EMAIL;TYPE=WORK:{lead['contact_email']}")

            # Phone
            if lead.get('phone'):
                vcf_lines.append(f"TEL;TYPE=WORK:{lead['phone']}")
            elif lead.get('contact_phone'):
                vcf_lines.append(f"TEL;TYPE=WORK:{lead['contact_phone']}")

            # Website
            if lead.get('company_website'):
                vcf_lines.append(f"URL:{lead['company_website']}")

            # Address
            if lead.get('address'):
                vcf_lines.append(f"ADR;TYPE=WORK:;;{lead.get('address', '')};{lead.get('city', '')};{lead.get('state', '')};{lead.get('postal_code', '')};{lead.get('country', '')}")

            vcf_lines.append("END:VCARD")
            vcf_lines.append("")  # Empty line between cards

        return "\n".join(vcf_lines).encode('utf-8')


class ExportManager:
    """
    Manages export operations for different formats
    """

    def __init__(self):
        self.exporters = {
            'csv': CSVExporter(),
            'excel': ExcelExporter(),
            'json': JSONExporter(),
            'xml': XMLExporter(),
            'vcf': VCFExporter(),
        }

    def export(
        self,
        leads: List[Dict[str, Any]],
        format: str,
        fields: Optional[List[str]] = None
    ) -> bytes:
        """
        Export leads in specified format

        Args:
            leads: List of leads
            format: Export format (csv, excel, json, xml, vcf)
            fields: Optional list of fields to include

        Returns:
            Exported data as bytes

        Raises:
            ValueError: If format is not supported
        """
        format = format.lower()

        if format not in self.exporters:
            raise ValueError(f"Unsupported export format: {format}")

        logger.info(f"Exporting {len(leads)} leads to {format} format")

        exporter = self.exporters[format]
        data = exporter.export(leads, fields)

        logger.info(f"Export completed: {len(data)} bytes")

        return data

    def get_supported_formats(self) -> List[str]:
        """Get list of supported export formats"""
        return list(self.exporters.keys())

    def get_file_extension(self, format: str) -> str:
        """
        Get file extension for format

        Args:
            format: Export format

        Returns:
            File extension with dot
        """
        extensions = {
            'csv': '.csv',
            'excel': '.xlsx',
            'json': '.json',
            'xml': '.xml',
            'vcf': '.vcf',
        }
        return extensions.get(format.lower(), '.txt')

    def get_mime_type(self, format: str) -> str:
        """
        Get MIME type for format

        Args:
            format: Export format

        Returns:
            MIME type string
        """
        mime_types = {
            'csv': 'text/csv',
            'excel': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'json': 'application/json',
            'xml': 'application/xml',
            'vcf': 'text/vcard',
        }
        return mime_types.get(format.lower(), 'application/octet-stream')


# Global export manager instance
export_manager = ExportManager()
